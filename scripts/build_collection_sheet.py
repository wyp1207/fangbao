#!/usr/bin/env python3
"""Build the on-site JCROB data collection workbook.

The workbook is a field checklist, not a replacement for the protocol PDF.
It intentionally leaves result columns blank for the engineer to fill on site.
"""

import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


OUT = r"D:\Desktop\防爆机器人\真机联调数据采集表.xlsx"

NAVY = "0B1F33"
HEADER = "123B5D"
CYAN = "00B4FF"
LIGHT = "EAF6FF"
WARN = "FFF3CD"
WHITE = "FFFFFF"
GRID = "2A4A66"

TITLE_FONT = Font(name="Microsoft YaHei", size=16, bold=True, color=WHITE)
SUB_FONT = Font(name="Microsoft YaHei", size=10, color="B9D9F2")
HEAD_FONT = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
BODY_FONT = Font(name="Microsoft YaHei", size=10, color="1F2937")
NOTE_FONT = Font(name="Microsoft YaHei", size=10, color="7A4B00")

thin = Side(style="thin", color=GRID)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def add_table(ws, title, note, headers, rows, widths, status_col=None, status_values=None, freeze="A4"):
    ncols = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = TITLE_FONT
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(row=2, column=1, value=note)
    c.font = SUB_FONT
    c.fill = PatternFill("solid", fgColor=HEADER)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 42

    for col, name in enumerate(headers, start=1):
        c = ws.cell(row=3, column=col, value=name)
        c.font = HEAD_FONT
        c.fill = PatternFill("solid", fgColor=CYAN)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[3].height = 30

    for r, row in enumerate(rows, start=4):
        for col, value in enumerate(row, start=1):
            c = ws.cell(row=r, column=col, value=value)
            c.font = BODY_FONT
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BORDER
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor=LIGHT)
        lines = 1
        for col, value in enumerate(row, start=1):
            if value is None:
                continue
            usable = max(8, widths[col - 1] * 1.7)
            lines = max(lines, math.ceil(len(str(value)) / usable))
        ws.row_dimensions[r].height = min(120, max(38, lines * 15))

    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A3:{get_column_letter(ncols)}{3 + len(rows)}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    if status_col:
        dv = DataValidation(type="list", formula1='"' + ",".join(status_values) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        col = get_column_letter(status_col)
        dv.add(f"{col}4:{col}{3 + len(rows)}")


def add_intro(wb):
    ws = wb.active
    ws.title = "00_使用说明"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 110

    ws.merge_cells("A1:B1")
    c = ws["A1"]
    c.value = "防爆机器人平台 · 真机联调数据采集表"
    c.font = Font(name="Microsoft YaHei", size=18, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 44

    rows = [
        ("版本", "v1.1 · 2026-10-20（按实机联调口径修订：move to index 原文带空格、is_online 不等于平台在线、热成像/云台后置、config 只读脱敏）"),
    ("采集对象", "JCROB 视觉/雷达导航轮式防爆机器人（本期）；海康可见光/热成像云台后续接入"),
    ("用途", "现场按表采集，回传后锁定 JCROB Adapter 的字段、单位和状态机；不是协议 PDF 的替代品。"),
        ("安全边界", "默认只读。任务下发已口头允许（2026-10-20），但现场仍需负责人签字；车体遥控、IO、充电必须单独确认。任何运动测试必须有安全员、急停手段和空旷区域。"),
        ("授权签字", "01 表第 15–17 行是权限档位，不是待测项。填“允许”= 签字人对该档写操作负责，必须写清：姓名 + 日期 + 允许范围 + 当时安全条件。平台只认这一栏，不认口头。"),
        ("填写方式", "每行填“现场填写/采集结果”，状态列选择 OK / 缺失 / 异常 / 未测。原始 JSON 存文件，表里只写相对路径或截图路径。"),
        ("敏感信息", "token、摄像头密码、Wi-Fi 密码不要写进本表；只记录“凭据存放位置”。probe_jcrob.py 已在落盘前自动脱敏。"),
        ("优先级", "P0 = 不做就无法开发；P1 = 影响体验或二期；P2 = 可后置/外部依赖。"),
        ("回传材料", "本 xlsx + 原始 JSON + 当前地图 PNG + 可见光/热成像截图 + 任务状态变化记录。"),
        ("关于 api.txt", "api.txt 是豆包/GLM 的模型密钥文件，不是机器人 API。机器人 API 是 PDF 里的 HTTP/WS；平台 API 是我们自己设计的 REST/WS。"),
        ("关于型号", "型号不是开发前提。需要的是 SDK/协议版本、固件版本、slam_type 和能力探测结果；平台按能力降级，不按型号写死。"),
        ("关于接口原文", "URL 按 PDF 原文，不做规范化。/navigation/move to index 原文带空格，请求时编码为 %20；只有原文路径返回 404/405，才启用别名 move_to_index。详见 docs/JCROB实机兼容性.md。"),
    ("海康视频", "本期后置；后续再接海康云台、可见光+热成像，不做 7×24 录像。"),
    ]
    r = 3
    for key, value in rows:
        a = ws.cell(row=r, column=1, value=key)
        b = ws.cell(row=r, column=2, value=value)
        a.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        a.fill = PatternFill("solid", fgColor=HEADER)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        a.border = BORDER
        b.font = BODY_FONT
        b.alignment = Alignment(vertical="top", wrap_text=True)
        b.border = BORDER
        ws.row_dimensions[r].height = 52
        r += 1

    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=2)
    c = ws.cell(row=r + 1, column=1, value="建议顺序：先跑 scripts/probe_jcrob.py 自动采集只读接口，再按本表补任务下发和现场问题；05 视频表等视频阶段再填。")
    c.font = NOTE_FONT
    c.fill = PatternFill("solid", fgColor=WARN)
    c.alignment = Alignment(wrap_text=True, vertical="center")


def build():
    wb = Workbook()
    add_intro(wb)

    ws = wb.create_sheet("01_现场基础信息")
    add_table(
        ws,
        "01 现场基础信息",
        "先把“能不能连、连到什么、允许做什么”确认清楚；这一页不通过，不要往下做写操作。",
        ["序号", "类别", "项目", "期望/说明", "现场填写", "备注"],
        [
            [1, "项目", "场站/项目名称", "例如：XX 化工厂 1 号站", "", ""],
            [2, "机器人", "机器人名称/编号", "平台展示名 + 唯一编号", "", ""],
            [3, "机器人", "机器人 IP", "固定 IP；记录网段/网关", "", ""],
            [4, "网络", "HTTP 端口", "默认 3546，前缀 /api/v1", "", ""],
            [5, "网络", "WebSocket 端口", "默认 3547", "", ""],
            [6, "网络", "UDP 发现端口", "默认 22002", "", ""],
            [7, "版本", "SDK/协议版本", "有 SDK 记 SDK 版本；没有就记 PDF/协议版本", "", ""],
            [8, "版本", "固件版本", "/system/info 的 info.version", "", ""],
            [9, "版本", "code_name", "/system/info 的 info.code_name", "", ""],
            [10, "身份", "机器人 ID / MAC", "/system/info 的 info.id / info.mac", "", ""],
            [11, "导航", "slam_type", "camera 或 lidar", "", ""],
            [12, "认证", "token 获取方式", "token / 用户名密码 / no_token_check", "", ""],
            [13, "地图", "当前地图名称", "/navigation/current_map 的 name", "", ""],
            [14, "地图", "当前楼层/level", "多楼层时必填", "", ""],
            [15, "权限", "任务下发", "需现场负责人签字：姓名 + 日期 + 允许范围。对话中已口头允许（2026-10-20），现场仍要落字", "", ""],
            [16, "权限", "车体遥控", "需签字；本期建议“不允许”，只做任务下发", "", "遥控绕过任务状态机，默认关闭"],
            [17, "权限", "IO/充电/对接", "需签字；首轮建议不测 IO，充电单独确认", "", "写操作不可重试，见方案 §5.7"],
            [18, "安全", "现场急停方式", "硬件急停位置 + 安全员", "", ""],
            [19, "视频", "海康云台 IP/型号", "记录型号、通道数、固件", "", ""],
            [20, "视频", "可见光通道", "通道号/主码流/子码流", "", ""],
            [21, "视频", "热成像通道", "通道号/主码流/子码流", "", ""],
            [22, "视频", "ONVIF/ISAPI", "是否支持 ONVIF；测温是否支持 ISAPI", "", ""],
            [23, "视频", "录像策略", "不做 7×24；记录事件/手动/告警片段", "", ""],
            [24, "网络", "平台到机器人链路", "固定 IP、路由、防火墙、是否跨网段", "", ""],
            [25, "记录", "采集人/日期", "姓名 + 日期 + 联系方式", "", ""],
        ],
        [6, 10, 22, 60, 44, 30],
        status_col=None,
    )

    ws = wb.create_sheet("02_HTTP必采数据")
    http_rows = [
        [1, "系统", "/system/status", "GET", "status", "Mapping / Navigating / Busy / Free", "状态互斥、任务前置检查", "", "未测", "P0"],
        [2, "系统", "/system/info", "GET", "info.version / code_name / id / mac / port", "字符串/端口", "版本与身份", "", "未测", "P0"],
        [3, "系统", "/system/info", "GET", "battery / camera_rgb / camera_depth / odom / imu / driver_port / charge / loop / slam_type", "百分比、Hz、bool、camera|lidar", "能力与健康", "", "未测", "P0"],
        [4, "系统", "/system/self_test", "GET", "motor_driver / camera / charge / lidar / bluetooth / battery", "bool/状态", "自检与告警", "", "未测", "P0"],
        [5, "系统", "/system/galileo_status", "GET", "navStatus / visualStatus / mapStatus / gcStatus / gbaStatus / chargeStatus / loopStatus / power / targetNumID / targetStatus / targetDistance / angleGoalStatus / controlSpeedX/Theta / currentSpeedX/Theta / currentPosX/Y / currentAngle / busyStatus", "状态、m、m/s、rad/s；currentAngle 单位用 --angle-check 多帧判定", "实时位姿/导航/充电/速度", "", "未测", "P0"],
        [6, "系统", "/system/config", "GET", "default_map / maps / auto_charge / slam_type / lf_linear_v / lf_theta_max", "配置（只读）", "默认地图、限速、自动充电；Adapter 只保留白名单字段，wifi_password 等不落盘不过前端", "", "未测", "P0"],
        [7, "地图", "/navigation/saved_maps", "GET", "resolution / origin / name / md5sum / map_type / levels / slam_type", "m/px、m、楼层", "地图列表与坐标换算", "", "未测", "P0"],
        [8, "地图", "/navigation/current_map", "GET", "name / md5sum", "字符串", "当前地图", "", "未测", "P0"],
        [9, "地图", "/navigation/map_png?name=", "GET", "PNG 二进制", "图片", "地图底图", "", "未测", "P0"],
        [10, "位姿", "/navigation/pose", "GET", "x / y / angle", "m / m / 度", "机器人位姿", "", "未测", "P0"],
        [11, "位姿", "/map/pose", "GET", "x / y / angle", "m / m / 度", "地图位姿校验", "", "未测", "P0"],
        [12, "导航", "/navigation/current_path", "GET", "map / path", "字符串", "当前地图/路径", "", "未测", "P0"],
        [13, "导航", "/navigation/nav_points", "GET", "map_name / path_name / path_data", "x y z theta errorX errorY direction name", "导航点/巡检点", "", "未测", "P0"],
        [14, "导航", "/navigation/path", "GET", "path_data", "x y z（每行）", "路径形状", "", "未测", "P1"],
        [15, "导航", "/navigation/point_connections", "GET", "points / connections", "坐标/连接关系", "跨地图/楼层连接", "", "未测", "P1"],
        [16, "巡航", "/navigation/loop_task", "GET", "循环任务对象/状态", "JSON", "读取当前循环巡航", "", "未测", "P0"],
        [17, "充电", "/navigation/charge_pose", "GET", "x / y / angle", "m / m / 度", "充电桩位姿/初始定位", "", "未测", "P0"],
        [18, "导航", "/navigation/target_distance", "GET", "distance", "m", "剩余距离", "", "未测", "P1"],
        [19, "导航", "/navigation/robot_tracks", "GET", "trajectory / track", "坐标序列", "历史轨迹", "", "未测", "P1"],
        [20, "导航", "/navigation/landmarks", "GET", "landmarks", "地图/二维码地标", "重定位辅助", "", "未测", "P1"],
        [21, "导航", "/navigation/virtual_wall", "GET", "虚拟墙数据", "字符串/坐标", "地图限制区", "", "未测", "P1"],
        [22, "导航", "/navigation/skip_area", "GET", "忽略区数据", "字符串/坐标", "地图限制区", "", "未测", "P1"],
        [23, "环境", "/navigation/wind_info", "GET", "current_speed / max_speed / level / direction", "m/s、级、方向", "风速风向", "", "未测", "P1"],
        [24, "任务", "/task", "GET", "任务列表/单个任务/state/sub_tasks/progress", "JSON", "任务模型与状态", "", "未测", "P0"],
        [25, "任务", "/action", "GET", "动作类型/参数 schema", "JSON", "动作能力清单", "", "未测", "P0"],
        [26, "任务", "/blockly_task", "GET", "title / description / is_enabled / blocklyxml", "JSON", "自动化任务", "", "未测", "P1"],
            [27, "视频", "/system/onvif_cameras", "GET", "摄像头列表/IP/名称/通道", "JSON", "发现可见光/热成像", "", "未测", "后置"],
            [28, "视频", "/system/onvif_camera/stream", "GET", "RTSP URL", "rtsp://", "视频转流输入", "", "未测", "后置"],
        [29, "云台", "/system/onvif_camera/pose", "GET", "x / y / zoom", "度/倍率", "云台位姿（本期后置）", "", "未测", "后置"],
        [30, "热成像", "/system/onvif_camera/temperature", "GET", "测温规则/最高/最低/平均/坐标", "℃", "热成像测温（本期后置，需摄像头支持红外测温 + 海康 ISAPI）", "", "未测", "后置"],
        [31, "IO", "/system/io", "GET", "out1..3", "0/1", "IO 状态", "", "未测", "P1"],
        [32, "IO", "/system/bw_io", "GET", "in/out 0..7", "0/1", "扩展板状态", "", "未测", "P1"],
        [33, "文件", "/system/saved_images", "GET", "文件列表", "路径/时间", "抓拍文件", "", "未测", "P1"],
        [34, "文件", "/system/saved_videos", "GET", "文件列表", "路径/时间", "录像文件", "", "未测", "P1"],
        [35, "网络", "/system/wifi", "GET", "SSID/信号/状态", "JSON", "网络健康", "", "未测", "P1"],
        [36, "网络", "/system/is_online", "GET", "能否连上互联网（result）", "bool", "机器人公网链路诊断；不是平台在线判定", "", "未测", "P1"],
        [37, "音频", "/system/volume", "GET", "音量", "0-100", "语音播报", "", "未测", "P1"],
        [38, "音频", "/system/tts/playing", "GET", "播放状态", "bool/JSON", "语音状态", "", "未测", "P1"],
        [39, "音频", "/system/tts/records", "GET", "播放记录", "JSON", "语音审计", "", "未测", "P1"],
    ]
    add_table(
        ws,
        "02 HTTP 必采数据",
        "先跑探针自动采集；表里填“是否拿到/字段是否一致/异常现象”。状态列建议统一用 OK/缺失/异常/未测。",
        ["序号", "类别", "接口", "方法", "关键字段", "期望/单位", "用途", "采集结果/文件路径", "状态", "优先级"],
        http_rows,
        [6, 10, 30, 8, 52, 34, 30, 40, 10, 8],
        status_col=9,
        status_values=["OK", "缺失", "异常", "未测"],
    )

    ws = wb.create_sheet("03_WebSocket样本")
    add_table(
        ws,
        "03 WebSocket 样本",
        "连接后每秒发送 ping；每个 topic 至少采 30 秒，记录频率、字段、断线表现。",
        ["序号", "Topic", "订阅地址", "关键字段", "期望单位", "采样时长", "样本文件", "实测频率", "备注"],
        [
            [1, "GalileoStatus", "ws://<ip>:3547/galileo/status", "位姿/速度/导航状态（schema 待确认）", "m、m/s、rad/s、度/弧度待确认", "30s", "", "", "是否与 /system/galileo_status 同构"],
            [2, "EnvSensorData", "ws://<ip>:3547/bw_env_sensors/EnvSensorData", "temperature / rh / smoke / pm1_0 / pm2_5 / pm10 / lel / noise", "℃、%RH、ppm、μg/m³、dB", "30s", "", "", "可燃气体单位与量程"],
            [3, "AirSensorData", "ws://<ip>:3547/bw_env_sensors/AirSensorData", "temperature / rh / co2 / illuminance", "℃、%RH、ppm、lux", "30s", "", "", "与 EnvSensorData 是否重复"],
            [4, "任务/充电/急停事件", "未知，需问厂家或抓包枚举", "task status / charge / estop", "状态枚举、时间戳", "30s", "", "", "若无推送，平台用 2–3s 轮询兜底"],
        ],
        [6, 18, 48, 54, 34, 12, 36, 12, 30],
    )

    ws = wb.create_sheet("04_任务与能力探测")
    add_table(
        ws,
        "04 任务与能力探测",
        "你已确认允许下发任务（现场仍需签字）。先测“单点导航 + 等待 + 拍照”，不要一上来测 IO/充电/对接。PDF 中不少写操作是 GET，平台内部必须改成 POST + 二次确认。已有复杂任务（行为树）读取要容错，但平台只产出扁平 action。",
        ["序号", "能力", "接口", "方法", "测试步骤", "期望结果", "实际结果", "可用性", "平台处理"],
        [
            [1, "任务创建", "/task", "POST", "创建一个只含 nav_action + sleep_action 的任务", "返回任务对象/ID", "", "待确认", "任务模板编译器"],
            [2, "任务启动", "/task/start", "GET（写）", "启动上一步任务，记录状态变化", "WAITTING→WORKING", "", "待确认", "平台 POST /tasks/{id}/start"],
            [3, "任务暂停", "/task/pause", "GET（写）", "运行中暂停 10s", "PAUSED", "", "待确认", "二次确认 + 审计"],
            [4, "任务继续", "/task/resume", "GET（写）", "继续到完成", "WORKING→COMPLETE", "", "待确认", "状态机"],
            [5, "任务停止", "/task/stop", "GET（写）", "另起任务后立即停止", "CANCELLED/停止原因", "", "待确认", "必须带 jcrob_task_id；不带 id 会取消全部任务，单列危险操作；≠ 安全急停"],
            [6, "任务循环", "/task/loop", "GET（写）", "读取 loop 参数并测试一次循环", "循环执行/状态", "", "待确认", "与 loop_task 二选一"],
            [7, "循环巡航", "/navigation/loop_task", "GET/POST/DELETE", "读取现状；如需测试，先建 2 点路线", "循环任务对象", "", "待确认", "固定路线优先"],
            [8, "导航启动", "/navigation/start", "GET（写）", "加载地图/路径并启动", "Navigating", "", "待确认", "首轮不测，任务接口优先"],
            [9, "保存的导航任务", "/navigation/start_nav_task", "POST", "启动已保存任务", "任务对象", "", "待确认", "可作为固定路线入口"],
            [10, "停止导航", "/navigation/stop_nav_task", "GET（写）", "停止当前导航", "停止/状态", "", "待确认", "权限 + 审计；停止任务 ≠ 停止导航，两者分开建模"],
            [11, "到点", "/navigation/move to index", "GET（写）", "移动到 1 号导航点（请求编码为 %20）", "导航任务对象", "", "待确认", "原文带空格，不做规范化；仅 404/405 才启用别名 move_to_index"],
            [12, "自动充电", "/navigation/go_charge", "GET（写）", "低电量时回充", "充电任务对象", "", "待确认", "先确认现场允许"],
            [13, "停止充电", "/navigation/stop_charge", "GET（写）", "停止充电任务", "任务对象", "", "待确认", "权限 + 审计"],
            [14, "IO 输出", "/system/io", "POST", "只读先看 out1..3；写操作待确认", "电平变化", "", "待确认", "默认禁用"],
            [15, "扩展 IO", "/system/bw_io", "PUT", "只读先看 in/out 0..7", "电平变化", "", "待确认", "默认禁用"],
            [16, "云台移动", "/system/onvif_camera/move", "GET（写）", "小角度移动后回正", "画面移动/位姿变化", "", "后置", "本期后置，不在联调范围"],
            [17, "云台预置位", "/system/onvif_camera/goto_preset", "GET（写）", "调用预置位 1", "画面到预置位", "", "后置", "本期后置，不在联调范围"],
            [18, "云台拍照", "/system/onvif_camera/snapshot", "GET（写）", "抓拍可见光/热成像", "返回图片路径", "", "后置", "本期后置，不在联调范围"],
            [19, "语音播报", "/action（tts_action）", "POST", "播报一句测试语音", "播放成功", "", "待确认", "音量/时段限制"],
            [20, "Blockly 自动化", "/blockly_task", "GET/POST", "读取现有任务；暂不新建", "blocklyxml", "", "待确认", "P1，先只读"],
            [21, "软件急停", "无（待查）", "—", "问厂家是否有 HTTP 急停/暂停接口；在空场地验证", "返回 ok，或明确不支持", "", "待确认", "不支持则平台“停止”=取消任务，页面必须标注"],
            [22, "通信超时停车", "无（待查）", "—", "断开平台网络 N 秒，观察车体是否自治停车", "车端在约定时间内停车", "", "待确认", "平台不做“失联后远程急停”的承诺"],
            [23, "热成像测温", "/system/onvif_camera/temperature", "GET", "对准温差明显物体，读取测温规则", "最高/最低/平均温度", "", "后置", "本期后置；后续按能力显示测温卡"],
            [24, "嵌套任务/行为树", "/task（sub_tasks 含 sequence/parallel/retry 等）", "GET", "读取机器人上已有的复杂任务，确认结构", "保留原始 JSON，页面降级只读", "", "待确认", "PDF p101–105 已定义；平台只产出扁平 action，读取必须容错，不能报接口失败"],
        ],
        [6, 16, 34, 16, 44, 34, 40, 10, 30],
        status_col=8,
        status_values=["是", "否", "待确认", "后置"],
    )

    ws = wb.create_sheet("05_视频与云台")
    add_table(
        ws,
        "05 视频与云台（本期后置，等明确型号再填）",
        "本期平台先不接视频/云台。后续确认两路 RTSP、ONVIF/ISAPI 测温、云台控制和延迟；录像只做任务抓拍/告警片段/手动。",
        ["序号", "项目", "采集内容", "期望/说明", "现场填写", "状态"],
        [
            [1, "设备", "海康云台型号", "铭牌型号/序列号/固件", "", "未测"],
            [2, "网络", "摄像头 IP/端口", "固定 IP；HTTP/RTSP/ONVIF 端口", "", "未测"],
            [3, "发现", "机器人 ONVIF 列表", "GET /system/onvif_cameras 是否发现两路", "", "未测"],
            [4, "可见光", "RTSP 主码流", "H.264/H.265、分辨率、帧率、码率", "", "未测"],
            [5, "可见光", "RTSP 子码流", "用于低带宽预览", "", "未测"],
            [6, "热成像", "RTSP 主码流", "分辨率、帧率、码率、伪彩", "", "未测"],
            [7, "热成像", "RTSP 子码流", "用于低带宽预览", "", "未测"],
            [8, "热成像", "测温方式", "点/线/框；最高/最低/平均；发射率设置", "", "未测"],
            [9, "热成像", "测温范围/精度", "例如 -20~150℃ / ±2℃（以铭牌为准）", "", "未测"],
            [10, "ISAPI", "测温接口", "JCROB /system/onvif_camera/temperature 是否可用；否则直连 ISAPI", "", "未测"],
            [11, "云台", "PTZ 控制", "上下左右/速度/回正/预置位", "", "未测"],
            [12, "云台", "变倍", "可见光/热成像是否联动", "", "未测"],
            [13, "抓拍", "可见光抓拍", "分辨率、存储路径、文件名规则", "", "未测"],
            [14, "抓拍", "热成像抓拍", "是否带温度叠加/原始温度矩阵", "", "未测"],
            [15, "延迟", "WebRTC 延迟", "目标 <1s；记录实测", "", "未测"],
            [16, "延迟", "FLV/HLS 延迟", "目标 2-3s；记录实测", "", "未测"],
            [17, "编码", "浏览器兼容", "H.264 优先；H.265 需转码", "", "未测"],
            [18, "音频", "对讲", "是否支持双向音频；协议", "", "未测"],
            [19, "录像", "事件录像策略", "任务开始/告警触发/手动；保留天数", "", "未测"],
            [20, "存储", "录像存放", "平台 MinIO/本地磁盘或海康 NVR", "", "未测"],
        ],
        [6, 12, 24, 58, 44, 10],
        status_col=6,
        status_values=["OK", "缺失", "异常", "未测"],
    )

    ws = wb.create_sheet("06_外部接口占位")
    add_table(
        ws,
        "06 外部接口占位（没有的先定契约）",
        "原则：平台内部 API 先稳定，外部适配器后接；没有数据时 UI 显示“未接入”，不要显示 0 或假数据。详细 schema 见 docs/外部接口占位规范.md。",
        ["序号", "能力", "平台内部接口（建议）", "数据方向", "关键字段", "传输方式（待定）", "当前状态", "没有数据时 UI 行为", "负责人/文档"],
        [
            [1, "BMS 单体电压", "GET /api/v1/devices/{id}/bms", "读", "totalVoltage/current/soc/cells[15]/temps/cycles/protection", "Modbus/RS485/MQTT/HTTP", "待厂家提供", "显示“BMS 未接入”，隐藏单体曲线", ""],
            [2, "固定设备控制", "POST /api/v1/devices/{id}/commands", "写", "command: door/light/ac/charger; params", "Modbus/MQTT/HTTP", "待接入", "按钮置灰 + 未配置", ""],
            [3, "AI 巡检", "POST /api/v1/ai/inspect", "写/读", "imageUrl/algorithm/roi → alarms[]", "HTTP/消息队列", "待接入", "隐藏 AI 结果，保留人工确认", ""],
            [4, "气象站", "GET /api/v1/weather/current", "读", "temperature/rh/wind/rain", "Modbus/HTTP", "待接入", "显示“气象站未接入”", ""],
            [5, "人脸识别", "POST /api/v1/face/verify", "写/读", "image/faceLib → matched/personId", "HTTP", "不排期", "菜单隐藏", ""],
            [6, "门禁", "POST /api/v1/access/command", "写", "open/close/authorize", "HTTP/Modbus", "待接入", "按钮置灰", ""],
            [7, "对讲", "POST /api/v1/intercom/session", "双向", "WebRTC/SIP/GB28181", "WebRTC/SIP", "待确认", "隐藏对讲按钮", ""],
            [8, "消防/气体固定点", "GET /api/v1/fire/status", "读", "smoke/lel/temperature/alarm", "Modbus/HTTP", "待接入", "只显示机器人移动监测", ""],
            [9, "电梯/门禁联动", "POST /api/v1/elevator/command", "写", "floor/open/close", "厂商协议", "待确认", "跨楼层任务禁用", ""],
            [10, "UPS/供电", "GET /api/v1/power/ups", "读", "input/output/battery/load", "SNMP/Modbus", "待接入", "显示“未接入”", ""],
            [11, "NVR/GB28181", "POST /api/v1/media/register", "读/写", "deviceId/channel/stream", "GB28181/ONVIF", "待确认", "只用 RTSP 预览", ""],
            [12, "工单/ERP", "POST /api/v1/workorders", "写", "alarmId/assignee/status", "HTTP", "待接入", "告警只在本平台闭环", ""],
        ],
        [6, 16, 40, 10, 48, 22, 14, 34, 24],
        status_col=7,
        status_values=["待接入", "待确认", "已接入", "不排期"],
    )

    ws = wb.create_sheet("07_现场问题记录")
    add_table(
        ws,
        "07 现场问题记录",
        "现场一定会遇到文档没写的情况；每条都记录请求/返回摘要和临时处理，回传后统一修 Adapter。",
        ["序号", "时间", "现象", "接口/设备", "请求/返回摘要", "影响", "临时处理", "待厂家确认", "负责人"],
        [
            [1, "", "", "", "", "", "", "", ""],
            [2, "", "", "", "", "", "", "", ""],
            [3, "", "", "", "", "", "", "", ""],
            [4, "", "", "", "", "", "", "", ""],
            [5, "", "", "", "", "", "", "", ""],
        ],
        [6, 16, 32, 26, 48, 26, 30, 30, 12],
    )

    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
