"""任务报告导出：xlsx 能打开、有任务信息与点位表。"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.db.session import SessionLocal
from app.models.task_run import TaskRun, TaskRunStep


def _make_run() -> int:
    with SessionLocal() as db:
        run = TaskRun(name="报告测试任务", status="done")
        db.add(run)
        db.commit()
        db.refresh(run)
        db.add(TaskRunStep(run_id=run.id, seq=1, point_name="P1", action="photo", status="ok"))
        db.add(TaskRunStep(run_id=run.id, seq=2, point_name="P2", action="photo",
                           status="failed", error="导航超时"))
        db.commit()
        return int(run.id)


def test_report_xlsx(auth_client):
    rid = _make_run()
    r = auth_client.get(f"/api/v1/task-runs/{rid}/report.xlsx")
    assert r.status_code == 200, r.text
    assert r.content[:2] == b"PK", "不是 xlsx（zip 头不对）"
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert ws.title == "巡检报告"
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "巡检报告" in text
    assert "P1" in text and "P2" in text
    assert "导航超时" in text
    assert "点位执行结果" in text
    assert "任务期间告警" in text
